"""Exportação de ativos para planilha (.xlsx): filtro, contagens e montagem do arquivo."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.asset_store import AssetStore
from app.models.climate_asset import ClimateAsset
from app.models.fire_asset import FireAsset
from app.models.water_asset import WaterAsset
from app.services.assets import apply_store_filters, normalize_key

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIMEZONE = ZoneInfo("America/Fortaleza")

# chave -> (modelo, nome da aba). A ordem daqui é a ordem das abas.
ASSET_TYPES = {
    "climatization": (ClimateAsset, "Climatização"),
    "fire_safety": (FireAsset, "Incêndio"),
    "water": (WaterAsset, "Água"),
}

# (cabeçalho, origem, atributo, tipo). Origem: "store" (loja) ou "asset" (equipamento).
_BEFORE = [
    ("Loja", "store", "store_name", "text"), ("BPCS", "store", "bpcs_number", "text"),
    ("SAP", "store", "sap_number", "text"), ("Praça", "store", "praca", "text"),
    ("Código", "asset", "asset_code", "text"), ("Tipo", "asset", "equipment_type", "text"),
    ("Local", "asset", "location", "text"),
]
_SPECIFIC = {
    "climatization": [
        ("Capacidade (BTU)", "capacity_btu", "int"), ("Marca", "brand", "text"), ("Modelo", "model", "text"),
        ("Nº de série", "serial_number", "text"), ("Ano de fabricação", "manufacture_year", "int"),
        ("Ano de instalação", "installation_year", "int"), ("Tensão", "voltage", "text"),
        ("Gás refrigerante", "refrigerant_gas", "text"),
    ],
    "fire_safety": [
        ("Agente extintor", "extinguisher_agent", "text"), ("Capacidade", "capacity", "text"),
        ("Ano de fabricação", "manufacture_year", "int"), ("Vencimento", "expiration_date", "date"),
        ("Teste hidrostático", "hydrostatic_test_date", "date"),
    ],
    "water": [
        ("Marca", "brand", "text"), ("Modelo", "model", "text"), ("Nº de série", "serial_number", "text"),
        ("Ano de fabricação", "manufacture_year", "int"), ("Ano de instalação", "installation_year", "int"),
        ("Tensão", "voltage", "text"), ("Última troca de filtro", "last_filter_change", "date"),
        ("Próxima troca de filtro", "next_filter_change", "date"),
    ],
}
_AFTER = [("Status", "asset", "status", "text"), ("Observações", "asset", "notes", "text")]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="10183E")


def columns_for(type_key: str) -> list[tuple[str, str, str, str]]:
    return _BEFORE + [(header, "asset", attr, kind) for header, attr, kind in _SPECIFIC[type_key]] + _AFTER


def parse_types(values: list[str] | None) -> list[str]:
    """Tipos pedidos, na ordem fixa das abas. Ausente = todos."""
    if values is None:
        return list(ASSET_TYPES)
    cleaned = [value.strip() for value in values if value and value.strip()]
    if not cleaned:
        raise HTTPException(status_code=400, detail="Escolha ao menos um tipo de ativo.")
    invalid = [value for value in cleaned if value not in ASSET_TYPES]
    if invalid:
        raise HTTPException(status_code=422, detail=f"Tipo de ativo inválido: {invalid[0]}.")
    return [key for key in ASSET_TYPES if key in cleaned]


def select_stores(db: Session, *, pracas: list[str] | None, store_ids: list[int] | None) -> list[AssetStore]:
    query = apply_store_filters(select(AssetStore).order_by(AssetStore.store_name), db, pracas=pracas, store_ids=store_ids)
    return list(db.scalars(query).all())


def count_assets(db: Session, stores: list[AssetStore], types: list[str]) -> dict[str, int]:
    """Equipamentos por tipo nas lojas dadas (tipos não pedidos contam zero)."""
    ids = [store.id for store in stores]
    counts = {key: 0 for key in ASSET_TYPES}
    if not ids:
        return counts
    for key in types:
        model = ASSET_TYPES[key][0]
        counts[key] = db.scalar(select(func.count(model.id)).where(model.store_id.in_(ids))) or 0
    return counts


def export_filename() -> str:
    return f"ativos-{datetime.now(TIMEZONE):%Y-%m-%d}.xlsx"


def _write(ws, row: int, column: int, value, kind: str) -> None:
    if value is None or value == "":
        return
    cell = ws.cell(row=row, column=column)
    if kind == "date" and isinstance(value, (date, datetime)):
        cell.value = value
        cell.number_format = "DD/MM/YYYY"
    elif kind == "int" and isinstance(value, (int, float)):
        cell.value = value
    else:
        cell.value = str(value)
        cell.data_type = "s"   # sempre texto: o que começa com = + - @ nunca vira fórmula


def build_workbook(db: Session, stores: list[AssetStore], types: list[str]) -> bytes:
    by_id = {store.id: store for store in stores}
    workbook = Workbook()
    workbook.remove(workbook.active)
    for key in types:
        model, title = ASSET_TYPES[key]
        columns = columns_for(key)
        sheet = workbook.create_sheet(title)
        widths = []
        for index, (header, *_rest) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
            cell.alignment = Alignment(vertical="center")
            widths.append(len(header))
        assets = list(db.scalars(select(model).where(model.store_id.in_(list(by_id)))).all()) if by_id else []
        assets.sort(key=lambda asset: (normalize_key(by_id[asset.store_id].store_name), asset.asset_code))
        for row, asset in enumerate(assets, start=2):
            store = by_id[asset.store_id]
            for column, (_header, source, attr, kind) in enumerate(columns, start=1):
                value = getattr(store if source == "store" else asset, attr)
                _write(sheet, row, column, value, kind)
                if value is not None:
                    widths[column - 1] = max(widths[column - 1], len(str(value)))
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 50)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
