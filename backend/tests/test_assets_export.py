"""Exportação de ativos para .xlsx."""

from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models.asset_store import AssetStore
from app.models.auth import AuditLog
from app.models.climate_asset import ClimateAsset
from app.models.fire_asset import FireAsset
from app.models.water_asset import WaterAsset
from tests.api_harness import SESSION_COOKIE, TestSession, fastapi_app, reset_database, seed_session
from tests.store_factory import add_asset_store

CLIMATIZACAO = ["Loja", "BPCS", "SAP", "Praça", "Código", "Tipo", "Local", "Capacidade (BTU)", "Marca", "Modelo",
                "Nº de série", "Ano de fabricação", "Ano de instalação", "Tensão", "Gás refrigerante", "Status", "Observações"]
INCENDIO = ["Loja", "BPCS", "SAP", "Praça", "Código", "Tipo", "Local", "Agente extintor", "Capacidade",
            "Ano de fabricação", "Vencimento", "Teste hidrostático", "Status", "Observações"]
AGUA = ["Loja", "BPCS", "SAP", "Praça", "Código", "Tipo", "Local", "Marca", "Modelo", "Nº de série",
        "Ano de fabricação", "Ano de instalação", "Tensão", "Última troca de filtro", "Próxima troca de filtro",
        "Status", "Observações"]
FORMULA = '=HYPERLINK("http://exemplo.test","clique")'


@pytest.fixture(autouse=True)
def clean():
    reset_database()
    yield
    reset_database()


@pytest.fixture
def lojas() -> dict[str, int]:
    """Natal Shopping (4 equipamentos), Midway (1), Mossoró Centro (2) e Fortaleza Sul (nenhum)."""
    db = TestSession()
    ids = {}
    for center_id, (nome, praca, bpcs, sap) in enumerate([
        ("Natal Shopping", "Natal", "111", "A1"), ("Midway", "Natal", "112", "A2"),
        ("Mossoró Centro", "Mossoró", "113", "A3"), ("Fortaleza Sul", "Fortaleza", "114", "A4"),
    ], start=1):
        loja = add_asset_store(db, center_id, nome, praca=praca, bcps=bpcs, sap=sap)
        ids[nome] = loja.id

    def clima(loja, codigo, local, **extra):
        db.add(ClimateAsset(store_id=ids[loja], asset_code=codigo, equipment_type="Split", capacity_btu=18000,
                            location=local, **extra))

    clima("Natal Shopping", "CLI-0002", "Sala de vendas", brand="Springer", model="Inverter", serial_number="SN-2")
    clima("Natal Shopping", "CLI-0001", FORMULA, notes="+cmd|' /C calc'!A0")
    db.add(FireAsset(store_id=ids["Natal Shopping"], asset_code="INC-0001", equipment_type="Extintor",
                     extinguisher_agent="PQS", capacity="6 kg", location="Corredor",
                     expiration_date=date(2027, 5, 10), manufacture_year=2023))
    db.add(WaterAsset(store_id=ids["Natal Shopping"], asset_code="AGU-0001", equipment_type="Purificador",
                      location="Copa", last_filter_change=date(2026, 3, 1)))
    clima("Midway", "CLI-0003", "Caixa")
    clima("Mossoró Centro", "CLI-0004", "Estoque")
    db.add(WaterAsset(store_id=ids["Mossoró Centro"], asset_code="AGU-0002", equipment_type="Gelágua", location="Refeitório"))
    db.commit()
    db.close()
    return ids


def cliente(permissoes=("assets.view",)) -> TestClient:
    token, _ = seed_session(list(permissoes))
    client = TestClient(fastapi_app, raise_server_exceptions=False)
    client.cookies.set(SESSION_COOKIE, token)
    return client


def baixar(client=None, params=None):
    return (client or cliente()).get("/assets/export", params=params or [])


def planilha(resposta):
    assert resposta.status_code == 200, resposta.text
    return load_workbook(BytesIO(resposta.content))


def linhas(ws) -> list[list]:
    return [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]


def test_traz_uma_aba_por_tipo_na_ordem_fixa(lojas):
    assert planilha(baixar()).sheetnames == ["Climatização", "Incêndio", "Água"]
    escolhidos = [("asset_type", "water"), ("asset_type", "climatization")]
    assert planilha(baixar(params=escolhidos)).sheetnames == ["Climatização", "Água"]


def test_cabecalhos_de_cada_aba(lojas):
    wb = planilha(baixar())
    for nome, esperado in (("Climatização", CLIMATIZACAO), ("Incêndio", INCENDIO), ("Água", AGUA)):
        assert [c.value for c in wb[nome][1]] == esperado


def test_linhas_ordenadas_por_loja_e_codigo(lojas):
    ws = planilha(baixar())["Climatização"]
    assert [(r[0], r[4]) for r in linhas(ws)] == [
        ("Midway", "CLI-0003"), ("Mossoró Centro", "CLI-0004"), ("Natal Shopping", "CLI-0001"), ("Natal Shopping", "CLI-0002")]


def test_dados_da_loja_e_do_equipamento_na_mesma_linha(lojas):
    linha = next(r for r in linhas(planilha(baixar())["Climatização"]) if r[4] == "CLI-0002")
    assert linha[:8] == ["Natal Shopping", "111", "A1", "Natal", "CLI-0002", "Split", "Sala de vendas", 18000]
    assert linha[8:11] == ["Springer", "Inverter", "SN-2"]
    assert linha[15] == "Operacional"


def test_datas_sao_datas_e_numeros_sao_numeros(lojas):
    ws = planilha(baixar())["Incêndio"]
    cabecalho = [c.value for c in ws[1]]
    linha = ws[2]
    vencimento = linha[cabecalho.index("Vencimento")]
    assert isinstance(vencimento.value, datetime) and vencimento.value.date() == date(2027, 5, 10)
    assert vencimento.number_format == "DD/MM/YYYY"
    assert linha[cabecalho.index("Ano de fabricação")].value == 2023
    assert linha[cabecalho.index("Teste hidrostático")].value is None   # vazio = célula vazia


def test_texto_com_formula_vira_texto_e_nunca_formula(lojas):
    resposta = baixar()
    ws = planilha(resposta)["Climatização"]
    linha = next(r for r in ws.iter_rows(min_row=2) if r[4].value == "CLI-0001")
    assert linha[6].value == FORMULA and linha[6].data_type == "s"
    assert linha[16].value == "+cmd|' /C calc'!A0" and linha[16].data_type == "s"
    with zipfile.ZipFile(BytesIO(resposta.content)) as arquivo:
        xml = "".join(arquivo.read(nome).decode("utf-8") for nome in arquivo.namelist() if nome.startswith("xl/worksheets/"))
    assert "<f>" not in xml and "<f " not in xml


def test_aba_sem_equipamentos_sai_so_com_o_cabecalho(lojas):
    resposta = baixar(params=[("praca", "Fortaleza")])
    wb = planilha(resposta)
    assert linhas(wb["Climatização"]) == [] and [c.value for c in wb["Climatização"][1]] == CLIMATIZACAO


def test_cabecalho_fixo_filtro_automatico_e_largura(lojas):
    ws = planilha(baixar())["Climatização"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions
    assert ws["A1"].font.bold and ws.column_dimensions["A"].width >= 10


def test_nome_do_arquivo_e_tipo(lojas):
    resposta = baixar()
    assert resposta.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert re.fullmatch(r'attachment; filename="ativos-\d{4}-\d{2}-\d{2}\.xlsx"', resposta.headers["content-disposition"])


def test_filtros_de_praca_e_loja_valem_na_exportacao(lojas):
    por_praca = baixar(params=[("praca", "Natal"), ("praca", "Mossoró")])
    assert {r[0] for r in linhas(planilha(por_praca)["Climatização"])} == {"Midway", "Mossoró Centro", "Natal Shopping"}
    estreito = baixar(params=[("praca", "Natal"), ("store_id", lojas["Midway"])])
    assert {r[0] for r in linhas(planilha(estreito)["Climatização"])} == {"Midway"}


def test_sem_lojas_ou_sem_tipo_da_400_e_tipo_invalido_da_422(lojas):
    vazio = baixar(params=[("praca", "Natal"), ("store_id", lojas["Mossoró Centro"])])
    assert vazio.status_code == 400 and "loja" in vazio.json()["detail"].lower()
    sem_tipo = baixar(params=[("asset_type", "")])
    assert sem_tipo.status_code == 400 and "tipo" in sem_tipo.json()["detail"].lower()
    assert baixar(params=[("asset_type", "ar-condicionado")]).status_code == 422
    assert cliente().get("/assets/export/preview", params=[("asset_type", "outro")]).status_code == 422


def test_previa_conta_lojas_e_ativos_por_tipo(lojas):
    previa = cliente().get("/assets/export/preview").json()
    assert previa == {"stores": 4, "climatization": 4, "fire_safety": 1, "water": 2}
    natal = cliente().get("/assets/export/preview", params=[("praca", "Natal")]).json()
    assert natal == {"stores": 2, "climatization": 3, "fire_safety": 1, "water": 1}
    so_clima = cliente().get("/assets/export/preview", params=[("asset_type", "climatization")]).json()
    assert so_clima == {"stores": 4, "climatization": 4, "fire_safety": 0, "water": 0}


@pytest.mark.parametrize("params", [[], [("praca", "Natal")], [("praca", "Mossoró"), ("asset_type", "water")]])
def test_contagem_da_previa_e_igual_as_linhas_do_arquivo(lojas, params):
    previa = cliente().get("/assets/export/preview", params=params).json()
    wb = planilha(baixar(params=params))
    chaves = {"Climatização": "climatization", "Incêndio": "fire_safety", "Água": "water"}
    for nome in wb.sheetnames:
        assert len(linhas(wb[nome])) == previa[chaves[nome]]


def test_exportacao_fica_na_auditoria(lojas):
    baixar(params=[("praca", "Natal")])
    db = TestSession()
    try:
        registro = db.query(AuditLog).filter_by(action="ASSET_EXPORT").one()
        assert registro.entity_type == "asset_export"
        assert registro.new_values["stores"] == 2 and registro.new_values["assets"] == 5
        assert registro.new_values["types"] == ["climatization", "fire_safety", "water"]
        assert registro.new_values["pracas"] == ["Natal"]
    finally:
        db.close()


def test_previa_nao_grava_auditoria(lojas):
    cliente().get("/assets/export/preview")
    db = TestSession()
    try:
        assert db.query(AuditLog).filter_by(action="ASSET_EXPORT").count() == 0
    finally:
        db.close()


def test_exige_sessao_e_permissao(lojas):
    anonimo = TestClient(fastapi_app, raise_server_exceptions=False)
    assert anonimo.get("/assets/export").status_code == 401
    assert anonimo.get("/assets/export/preview").status_code == 401
    sem_permissao = cliente(["assistant.use"])
    assert sem_permissao.get("/assets/export").status_code == 403
    assert sem_permissao.get("/assets/export/preview").status_code == 403


def test_caracteres_de_controle_nao_derrubam_a_exportacao(lojas):
    db = TestSession()
    db.add(ClimateAsset(store_id=lojas["Midway"], asset_code="CLI-0009", equipment_type="Split",
                        location="Sala\x0b2", notes="linha\x00 colada\x1f do Word"))
    db.commit()
    db.close()
    ws = planilha(baixar())["Climatização"]
    linha = next(r for r in ws.iter_rows(min_row=2) if r[4].value == "CLI-0009")
    assert linha[6].value == "Sala2" and linha[16].value == "linha colada do Word"
